from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login as auth_login
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import AuthenticationForm
from django.urls import reverse
from django.utils import timezone
from django.contrib import messages
from django.db.models import Q
from apps.classes.models import Class, ClassSession
from apps.reservations.models import Reservation
from apps.users.models import User, Coach
from django.http import HttpResponse
from django.contrib.admin.views.decorators import staff_member_required
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


@login_required
def dashboard(request):
    hoy = timezone.now().date()

    clases_activas = Class.objects.filter(is_active=True).count()
    sesiones_hoy = ClassSession.objects.filter(
        start_datetime__date=hoy,
        status='SCHEDULED'
    ).count()
    reservas_totales = Reservation.objects.count()
    usuarios_registrados = User.objects.count()

    return render(request, 'dashboard.html', {
        'user': request.user,
        'clases_activas': clases_activas,
        'sesiones_hoy': sesiones_hoy,
        'reservas_totales': reservas_totales,
        'usuarios_registrados': usuarios_registrados,
    })

def home(request):
    return render(request, 'home.html', {
        'user': request.user,
    })

def login_view(request):
    if request.method == 'POST':
        form = AuthenticationForm(request, data=request.POST)
        if form.is_valid():
            user = form.get_user()
            if user.role == User.Roles.ADMIN:
                return redirect(f"{reverse('login')}?error=role_mismatch")
            auth_login(request, user)
            return redirect('dashboard')
        return redirect(f"{reverse('login')}?error=true")

    return render(request, 'registration/login.html', {
        'open_admin_modal': request.GET.get('admin') == '1'
    })


def admin_login(request):
    if request.method == 'GET':
        return render(request, 'registration/login.html', {
            'open_admin_modal': True,
        })

    form = AuthenticationForm(request, data=request.POST)
    if form.is_valid():
        user = form.get_user()
        if user.role != User.Roles.ADMIN:
            return redirect(f"{reverse('login')}?error=role_mismatch&admin=1")
        auth_login(request, user)
        return redirect('dashboard')

    return redirect(f"{reverse('login')}?error=true&admin=1")


def register(request):
    if request.method == 'POST':
        correo = request.POST.get('correo', '').strip().lower()
        nombre = request.POST.get('nombre', '').strip()
        apellido = request.POST.get('apellido', '').strip()
        telefono = request.POST.get('telefono', '').strip()
        password = request.POST.get('password', '')

        User.objects.create_user(
            username=nombre.lower() + apellido.lower(),
            email=correo,
            first_name=nombre,
            last_name=apellido,
            phone=telefono,
            password=password,
            role=User.Roles.CLIENT
        )

        messages.success(
            request,
            'Cuenta creada correctamente. Ya puedes iniciar sesión.'
        )

        return redirect('login')

    return render(request, 'registro.html')

def register_admin(request):
    if request.method == 'POST':
        correo = request.POST.get('correo', '').strip().lower()
        nombre = request.POST.get('nombre', '').strip()
        apellido = request.POST.get('apellido', '').strip()
        telefono = request.POST.get('telefono', '').strip()
        password = request.POST.get('password', '')

        User.objects.create_user(
            username=nombre.lower() + apellido.lower(),
            email=correo,
            first_name=nombre,
            last_name=apellido,
            phone=telefono,
            password=password,
            role=User.Roles.ADMIN
        )

        messages.success(
            request,
            'Cuenta de administrador creada correctamente. Ya puedes iniciar sesión.'
        )

        return redirect('admin_login')

    return render(request, 'registration/registroAdmin.html')

@login_required
def users_list(request):
    if request.user.role != 'ADMIN':
        return redirect('dashboard')
    
    username_filter = request.GET.get('username', '').strip()
    role_filter = request.GET.get('role', '')
    status_filter = request.GET.get('status', '')

    usuarios = User.objects.all().order_by('role', 'username')

    if username_filter:
        usuarios = usuarios.filter(
            Q(username__icontains=username_filter) |
            Q(first_name__icontains=username_filter) |
            Q(last_name__icontains=username_filter)
        )

    if role_filter in [User.Roles.ADMIN, User.Roles.COACH, User.Roles.CLIENT]:
        usuarios = usuarios.filter(role=role_filter)

    if status_filter == 'active':
        usuarios = usuarios.filter(is_active=True)
    elif status_filter == 'inactive':
        usuarios = usuarios.filter(is_active=False)

    return render(request, 'users.html', {
        'usuarios': usuarios,
        'user': request.user,
        'filters': {
            'username': username_filter,
            'role': role_filter,
            'status': status_filter,
        },
        'results_count': usuarios.count(),
    })


@login_required
def user_toggle(request, pk):
    if request.user.role != 'ADMIN':
        return redirect('dashboard')
    
    usuario = get_object_or_404(User, pk=pk)
    usuario.is_active = not usuario.is_active
    usuario.save()
    return redirect('users')

@login_required
def coaches_list(request):

    if request.user.role != User.Roles.ADMIN:
        return redirect('dashboard')

    name_filter = request.GET.get('name', '').strip()
    status_filter = request.GET.get('status', '')

    coaches = User.objects.filter(
        role=User.Roles.COACH
    ).order_by('username')

    if name_filter:
        coaches = coaches.filter(
            Q(username__icontains=name_filter) |
            Q(first_name__icontains=name_filter) |
            Q(last_name__icontains=name_filter)
        )

    if status_filter == 'active':
        coaches = coaches.filter(is_active=True)
    elif status_filter == 'inactive':
        coaches = coaches.filter(is_active=False)

    return render(
        request,
        'coaches.html',
        {
            'coaches': coaches,
            'filters': {
                'name': name_filter,
                'status': status_filter,
            },
            'results_count': coaches.count(),
        }
    )

@login_required
def coach_create(request):

    if request.user.role != User.Roles.ADMIN:
        return redirect('dashboard')

    if request.method == 'POST':

        User.objects.create_user(
            username=request.POST['username'],
            email=request.POST['email'],
            first_name=request.POST.get('first_name', ''),
            last_name=request.POST.get('last_name', ''),
            phone=request.POST.get('phone', ''),
            password=request.POST['password'],
            role=User.Roles.COACH
        )

        messages.success(
            request,
            'Entrenador creado correctamente.'
        )

        return redirect('coaches')

    coaches = User.objects.filter(
        role=User.Roles.COACH
    )

    return render(
        request,
        'coaches.html',
        {
            'coaches': coaches,
            'show_form': True
        }
    )


@login_required
def coach_toggle(request, pk):

    if request.user.role != User.Roles.ADMIN:
        return redirect('dashboard')

    coach = get_object_or_404(
        User,
        pk=pk,
        role=User.Roles.COACH
    )

    coach.is_active = not coach.is_active
    coach.save()

    return redirect('coaches')

@login_required
def exportar_usuarios_excel(request):
    """
    Exporta todos los usuarios a un archivo Excel
    Solo disponible para administradores
    """
    # Verificar que el usuario sea administrador
    if request.user.role != 'ADMIN':
        messages.error(request, 'No tienes permiso para exportar usuarios.')
        return redirect('dashboard')
    
    # Obtener todos los usuarios (puedes aplicar los mismos filtros que en users_list si quieres)
    usuarios = User.objects.all().order_by('role', 'username')
    
    # Aplicar los mismos filtros que en users_list (opcional)
    username_filter = request.GET.get('username', '').strip()
    role_filter = request.GET.get('role', '')
    status_filter = request.GET.get('status', '')

    if username_filter:
        usuarios = usuarios.filter(
            Q(username__icontains=username_filter) |
            Q(first_name__icontains=username_filter) |
            Q(last_name__icontains=username_filter)
        )

    if role_filter in [User.Roles.ADMIN, User.Roles.COACH, User.Roles.CLIENT]:
        usuarios = usuarios.filter(role=role_filter)

    if status_filter == 'active':
        usuarios = usuarios.filter(is_active=True)
    elif status_filter == 'inactive':
        usuarios = usuarios.filter(is_active=False)
    
    # Crear libro de trabajo
    wb = Workbook()
    ws = wb.active
    ws.title = "Usuarios"
    
    # Definir encabezados
    headers = [
        'ID', 'Usuario', 'Email', 'Nombre', 'Apellido', 
        'Rol', 'Teléfono', 'Fecha de registro', 'Último login', 'Activo', 'Staff'
    ]
    
    # Estilos para encabezados
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    
    # Escribir encabezados
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    # Escribir datos
    for row, user in enumerate(usuarios, 2):
        ws.cell(row=row, column=1, value=user.id)
        ws.cell(row=row, column=2, value=user.username)
        ws.cell(row=row, column=3, value=user.email or '')
        ws.cell(row=row, column=4, value=user.first_name or '')
        ws.cell(row=row, column=5, value=user.last_name or '')
        ws.cell(row=row, column=6, value=user.get_role_display())
        ws.cell(row=row, column=7, value=user.phone or '')
        ws.cell(row=row, column=8, value=user.date_joined.strftime('%d/%m/%Y %H:%M') if user.date_joined else '')
        ws.cell(row=row, column=9, value=user.last_login.strftime('%d/%m/%Y %H:%M') if user.last_login else 'Nunca')
        ws.cell(row=row, column=10, value='Sí' if user.is_active else 'No')
        ws.cell(row=row, column=11, value='Sí' if user.is_staff else 'No')
    
    # Ajustar ancho de columnas automáticamente
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 40)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Configurar respuesta HTTP para descarga
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=usuarios_nexus.xlsx'
    
    wb.save(response)
    return response