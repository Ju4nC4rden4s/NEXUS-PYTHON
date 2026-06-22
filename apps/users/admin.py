# apps/users/admin.py

from django.contrib import admin
from django.contrib.auth.admin import UserAdmin as BaseUserAdmin
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .models import User, Coach


def exportar_usuarios_excel(modeladmin, request, queryset):
    """
    Acción personalizada para exportar usuarios a Excel
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Usuarios"

    headers = [
        'ID', 'Usuario', 'Email', 'Nombre', 'Apellido', 
        'Rol', 'Teléfono', 'Fecha de registro', 'Último login', 'Activo', 'Staff'
    ]
    
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row, user in enumerate(queryset, 2):
        ws.cell(row=row, column=1, value=user.id)
        ws.cell(row=row, column=2, value=user.username)
        ws.cell(row=row, column=3, value=user.email or '')
        ws.cell(row=row, column=4, value=user.first_name or '')
        ws.cell(row=row, column=5, value=user.last_name or '')
        ws.cell(row=row, column=6, value=user.get_role_display())
        ws.cell(row=row, column=7, value=user.phone or '')
        ws.cell(row=row, column=8, value=user.date_joined.strftime('%Y-%m-%d %H:%M') if user.date_joined else '')
        ws.cell(row=row, column=9, value=user.last_login.strftime('%Y-%m-%d %H:%M') if user.last_login else 'Nunca')
        ws.cell(row=row, column=10, value='Sí' if user.is_active else 'No')
        ws.cell(row=row, column=11, value='Sí' if user.is_staff else 'No')

    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 30)
        ws.column_dimensions[column_letter].width = adjusted_width

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=usuarios_nexus.xlsx'
    
    wb.save(response)
    return response

exportar_usuarios_excel.short_description = "📊 Exportar usuarios seleccionados a Excel"


# SOLO UN REGISTRO para el modelo User
@admin.register(User)
class CustomUserAdmin(BaseUserAdmin):
    """
    Admin personalizado para el modelo User
    """
    list_display = ('username', 'email', 'first_name', 'last_name', 'role', 'phone', 'is_active', 'date_joined')
    list_filter = ('role', 'is_active', 'is_staff')
    search_fields = ('username', 'email', 'first_name', 'last_name', 'phone')
    actions = [exportar_usuarios_excel]
    
    # Heredamos los fieldsets de UserAdmin y agregamos 'role' y 'phone'
    fieldsets = BaseUserAdmin.fieldsets + (
        ('Información adicional', {'fields': ('role', 'phone')}),
    )
    
    add_fieldsets = BaseUserAdmin.add_fieldsets + (
        ('Información adicional', {'fields': ('role', 'phone')}),
    )


# SOLO UN REGISTRO para el modelo Coach (proxy)
@admin.register(Coach)
class CoachAdmin(admin.ModelAdmin):
    """
    Admin para el modelo proxy Coach (solo muestra usuarios con rol COACH)
    """
    list_display = ('username', 'email', 'first_name', 'last_name', 'phone', 'is_active')
    search_fields = ('username', 'email', 'first_name', 'last_name')
    
    def get_queryset(self, request):
        """Filtrar solo usuarios con rol COACH"""
        return super().get_queryset(request).filter(role=User.Roles.COACH)
    
    def has_add_permission(self, request):
        """Los coaches se crean desde el admin de usuarios"""
        return False